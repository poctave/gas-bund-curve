"""Assemble the research workbook: narrative sheets (summary, intuition, benchmarks, test catalog)
plus every result table from output/tables, charts, and the data dictionary.

Headline numbers in the narrative sheets are live cross-sheet formulas into the result sheets, so the
summary always matches the tables it cites. Result tables are values exported by econometrics.py."""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import CLEAN, FIGURES, ROOT, TABLES

OUT = ROOT / "output" / "Gas_Bund_Research_Workbook.xlsx"
FONT = "Arial"
NAVY, LIGHT, SIG = "1F3864", "EEF2F8", "E2F0D9"
H_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
H_FILL = PatternFill("solid", fgColor=NAVY)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
LINK = Font(name=FONT, size=10, color="008000")  # green = cross-sheet link
TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
NOTE = Font(name=FONT, size=9, italic=True, color="52514E")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))

wb = Workbook()
wb.remove(wb.active)
TABLE_POS: dict[str, tuple[int, list[str]]] = {}  # sheet -> (header row, columns)
FRAMES: dict[str, pd.DataFrame] = {}


# --------------------------------------------------------------------------- helpers
def fmt_for(col: str) -> str:
    c = col.lower()
    if c in ("n", "n_obs", "n_weeks", "n_months", "n_forecasts", "lags", "lag_order_bic", "lag_aic",
             "lag_hqic", "horizon", "horizon_days", "lag_months", "rank_at_5pct"):
        return "0"
    if c.endswith("_p") or c.startswith("p_") or "_p_" in c or c.endswith("pvalue") or c in ("chow_p", "granger_p", "wald_p", "clark_west_p", "pred_p", "cum_p"):
        return "0.000"
    if "share" in c or c in ("oos_r2", "r2", "pred_r2"):
        return "0.0%" if "share" in c else "0.000"
    return "0.00"


def is_p(col: str) -> bool:
    return fmt_for(col) == "0.000" and ("p" in col.lower())


def sheet_header(ws, title: str, notes: list[str]) -> int:
    ws["A1"] = title
    ws["A1"].font = TITLE
    r = 2
    for n in notes:
        ws.cell(row=r, column=1, value=n).font = NOTE
        r += 1
    return r + 1


def write_df(ws, df: pd.DataFrame, start: int, widths: dict[str, int] | None = None) -> int:
    for j, c in enumerate(df.columns, 1):
        cell = ws.cell(row=start, column=j, value=c)
        cell.font, cell.fill, cell.alignment = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.floating, float)) and (np.isnan(v) or np.isinf(v)):
                v = None
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, np.floating):
                v = float(v)
            elif isinstance(v, (np.bool_,)):
                v = bool(v)
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=v)
            cell.font = BODY
            if isinstance(v, float):
                cell.number_format = fmt_for(df.columns[j - 1])
            cell.border = THIN
    last = start + len(df)
    for j, c in enumerate(df.columns, 1):
        L = get_column_letter(j)
        w = (widths or {}).get(c)
        if w is None:
            sample = [len(str(x)) for x in df[c].head(50)] + [min(len(c), 22)]
            w = min(max(sample) + 2, 48)
        ws.column_dimensions[L].width = max(w, 9)
        if is_p(c) and len(df):
            ws.conditional_formatting.add(f"{L}{start + 1}:{L}{last}",
                                          CellIsRule(operator="lessThan", formula=["0.05"], fill=PatternFill("solid", fgColor=SIG)))
    ws.row_dimensions[start].height = 30
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    return last


def result_sheet(name: str, csv: str, title: str, notes: list[str], cols: list[str] | None = None,
                 rename: dict[str, str] | None = None) -> pd.DataFrame:
    df = pd.read_csv(TABLES / f"{csv}.csv")
    if cols:
        df = df[cols]
    ws = wb.create_sheet(name)
    start = sheet_header(ws, title, notes + [f"Source: src/econometrics.py -> output/tables/{csv}.csv. "
                                                 "Green cells: p < 0.05."])
    write_df(ws, df, start)
    TABLE_POS[name] = (start, list(df.columns))
    FRAMES[name] = df
    return df


def ref(sheet: str, mask, col: str) -> str:
    """Absolute cross-sheet reference to the cell in `sheet` for the first row matching mask."""
    start, cols = TABLE_POS[sheet]
    df = FRAMES[sheet]
    idx = np.flatnonzero(np.asarray(mask))
    if len(idx) == 0:
        raise KeyError(f"no row in {sheet} for {col}")
    r = start + 1 + int(idx[0])
    c = get_column_letter(cols.index(col) + 1)
    return f"='{sheet}'!${c}${r}"


def val(sheet: str, mask, col: str) -> float:
    df = FRAMES[sheet]
    return float(df.loc[np.asarray(mask), col].iloc[0])


# --------------------------------------------------------------------------- result sheets
def build_results() -> None:
    result_sheet("Descriptive", "01_descriptive_stats", "Descriptive statistics of daily changes",
                 ["Rates in bp/day, prices in %/day (log). Jarque-Bera: H0 normal. Ljung-Box Q(10): H0 no "
                  "autocorrelation. ARCH-LM(5): H0 no volatility clustering."])
    result_sheet("Unit_Roots", "02_unit_roots", "Stationarity: ADF, KPSS and Zivot-Andrews (weekly levels & changes)",
                 ["ADF / Zivot-Andrews H0: unit root. KPSS H0: stationary. 'stationary' = ADF rejects and KPSS does "
                  "not. Zivot-Andrews allows one endogenous break (break date reported)."])
    result_sheet("Cointegration", "03_cointegration", "Cointegration: Engle-Granger (pairs) and Johansen (systems)",
                 ["H0 (EG): no cointegration. Johansen: sequential trace test at 5%. Weekly levels."])
    result_sheet("Transmission_Daily", "04_transmission_daily",
                 "Transmission regressions, daily: change in rate on TTF, Brent, EURUSD - by regime",
                 ["bp_per_10pct_ttf = total effect of a +10% TTF move in bp (same day + previous day for Bundesbank "
                  "yields, which are snapshotted before the market close). Newey-West HAC SEs.",
                  "Multiple testing (Holm, within each regime): p_ttf_holm = across all 20 dependent variables (strictest); "
                  "p_ttf_holm_family = within its family (rates & spreads: 13 series; curve shape: 7 series)."],
                 cols=["dependent", "family", "regime", "n", "async_adjusted", "bp_per_10pct_ttf", "t_ttf", "p_ttf",
                       "p_ttf_holm", "p_ttf_holm_family",
                       "beta_ttf_same_day", "beta_ttf_prev_day", "bp_per_10pct_brent", "t_brent", "p_brent", "r2"])
    result_sheet("Transmission_Weekly", "04_transmission_weekly",
                 "Transmission regressions, weekly (Fri-Fri): robustness to data timing",
                 ["Gas/Brent weekly returns = sum of roll-cleaned daily returns. For Bundesbank yields the market "
                  "window is lagged one business day (Fri-Thu)."],
                 cols=["dependent", "family", "regime", "n", "bp_per_10pct_ttf", "t_ttf", "p_ttf", "p_ttf_holm",
                       "p_ttf_holm_family", "bp_per_10pct_brent", "t_brent", "p_brent", "r2"])
    result_sheet("Regime_Tests", "05_regime_interaction",
                 "Is the gas effect different in the war and in 2021-22? Pooled model with regime interactions",
                 ["Base = 2014-21 + 2023-Feb26 combined. H0: regime total effect = base total effect (HAC Wald)."])
    result_sheet("Diagnostics", "06_residual_diagnostics", "Residual diagnostics of the transmission regressions",
                 ["Durbin-Watson (~2 = no AR(1)); Breusch-Godfrey(5) & Ljung-Box(10): H0 no serial correlation; "
                  "Breusch-Pagan & White: H0 homoskedastic; ARCH-LM(5): H0 no ARCH; Jarque-Bera: H0 normal; "
                  "Ramsey RESET: H0 correct functional form; CUSUM: H0 stable coefficients; VIF > 5 = collinearity.",
                  "Here green = rejection of the null, i.e. a problem detected (handled by HAC SEs / regime split)."])
    result_sheet("Structural_Breaks", "07_structural_breaks", "Structural breaks in the gas sensitivity (2023 - 2026)",
                 ["Chow F at 2 Mar 2026 (all coefficients). Robust Wald: shift in the total TTF effect at the war date "
                  "(HAC). Quandt-Andrews sup-Wald: break date unknown, 15% trimming; 5% critical value 8.68 (Andrews 1993)."])
    result_sheet("Timing", "09_lead_lag_timing", "Data timing & lead/lag (Dimson) tests",
                 ["Rows 1-3: Bundesbank yields load equally on previous- and same-day LSEG OIS moves => yields are "
                  "observed ~half a day earlier. Remaining rows: TTF at t-1, t, t+1."])
    result_sheet("Mediation", "10_mediation", "Mediation: how much of the gas effect runs through inflation expectations and policy pricing?",
                 ["Daily rows: TTF -> inflation swap -> €STR OIS 1y1y (all LSEG closes, synchronous).",
                  "Weekly rows: full serial chain TTF -> ILS 1y1y -> OIS 1y1y -> Bund 2Y. CIs: moving-block bootstrap, 2,000 draws.",
                  "Effects in bp per 1% TTF (bp_per_10pct_ttf column x10%); path b / d21 coefficients in bp per bp."])
    result_sheet("VAR_Diagnostics", "11a_var_diagnostics", "VAR specification & diagnostics",
                 ["Lag order by BIC. Stable = all roots inside unit circle. Portmanteau H0: residual whiteness."])
    result_sheet("Granger", "11b_granger_causality", "Granger causality (VAR-based F tests)",
                 ["H0: lags of 'cause' do not help predict 'effect'. Daily VAR uses LSEG-synchronous series only."])
    result_sheet("VAR_FEVD", "11c_var_fevd", "Forecast-error variance decomposition (Cholesky, two orderings)",
                 ["Share of each variable's forecast-error variance due to each shock. 'Brent first' vs 'TTF first' "
                  "bounds the gas contribution when oil and gas shocks are correlated."])
    result_sheet("VAR_IRF", "11d_var_irf_ttf_shock", "Impulse responses to a 1-s.d. orthogonalised TTF shock",
                 ["Units: bp (rates) per 1-s.d. TTF shock."])
    result_sheet("Local_Projections", "12_local_projections", "Local projections: persistence and predictivity",
                 ["cum_*: y(t+h) - y(t-1) on gas(t) = impact and persistence. pred_*: y(t+h) - y(t+s) on gas(t) = does "
                  "gas forecast LATER moves? (s=1 for Bundesbank yields to skip the overlapping day). HAC lags h+1."])
    result_sheet("OOS_Forecast", "13_oos_forecast", "Out-of-sample forecasting of next week's move",
                 ["Expanding window. OOS R2 > 0 = beats benchmark. Clark-West: one-sided test for nested models "
                  "(H0 equal accuracy)."])
    result_sheet("HICP_Passthrough", "14_hicp_passthrough", "Gas and oil pass-through to realised euro-area HICP",
                 ["Model A: m/m HICP on 0-12 monthly lags of TTF and Brent (EUR) + month dummies; cumulative effect in pp of "
                  "price level per +10%.", "Model B: HICP y/y on 12-month energy price changes lagged L months (lag scan)."])
    result_sheet("Curve_PCA", "16_curve_pca", "Bund curve principal components and their gas sensitivity",
                 ["PC1 = level, PC2 = slope (+ = steepening), PC3 = curvature. Sensitivities in bp per +10% TTF."])
    result_sheet("Curve_Slopes", "17_curve_slopes", "Which flattener carries gas? All Bund slopes compared",
                 ["Slope = long yield minus short yield, bp. Sensitivities in bp of slope change per +10% TTF (negative = "
                  "flattening). Daily = same + previous day (timing-adjusted); weekly = Fri-Fri.",
                  "Lower rows: how each slope actually moved during the war. Tenor moves: second table below."])
    ws = wb["Curve_Slopes"]
    tm = pd.read_csv(TABLES / "17b_tenor_moves_war.csv")
    r0 = ws.max_row + 3
    ws.cell(row=r0 - 1, column=1, value="Bund yield change by tenor, 27 Feb - 9 Oct 2026").font = BOLD
    write_df(ws, tm, r0)
    ws.freeze_panes = ws.cell(row=TABLE_POS["Curve_Slopes"][0] + 1, column=2)
    result_sheet("OLS_Fit", "18b_ols_fit", "OLS transmission regressions: specification and fit statistics",
                 ["One row per regression. Daily: dy(t) = a + b g(t) + c o(t) + d e(t) [+ the same at t-1 for Bundesbank series] "
                  "+ u(t); g, o, e = 100 x log change of TTF, Brent, EURUSD; dy in bp. Weekly: Fri-Fri changes.",
                  "hac_lags = Newey-West lags floor(4 (T/100)^(2/9)). robust_F = HAC Wald test that all slopes are zero. "
                  "ttf_total = (b0 + b1) x 100 ln(1.1), bp per +10% TTF."])
    result_sheet("OLS_Coefficients", "18_ols_coefficients", "OLS transmission regressions: every coefficient",
                 ["Coefficients in bp per 1% move of the regressor; Newey-West HAC standard errors. Filter by frequency, "
                  "regime and dependent variable."])
    result_sheet("Rolling_Betas", "08_rolling_betas_120d", "Rolling 120-day TTF sensitivity (bp per +10%)",
                 ["Stepped weekly. 95% band from HAC SEs."])
    event_sheet()


def event_sheet() -> None:
    df = pd.read_csv(TABLES / "15_event_study.csv")
    ws = wb.create_sheet("Event_Study")
    start = sheet_header(ws, "Event study: key war and policy dates (levels; changes computed by formula)",
                         ["Bund yields are snapshotted before the close, so the 2-day change (t-1 -> t+1) is the comparable "
                          "window for them. Rates (%) converted to bp; spreads already in bp; prices as % change.",
                          "Source: output/tables/15_event_study.csv (levels). Change columns are live formulas."])
    df = df.copy()
    df["1-day change"] = ""
    df["2-day change"] = ""
    df["unit"] = ""
    write_df(ws, df, start)
    for i in range(len(df)):
        r = start + 1 + i
        var = df.loc[i, "variable"]
        if "(%)" in var:
            e1, e2, unit = f"(E{r}-D{r})*100", f"(F{r}-D{r})*100", "bp"
        elif "(bp)" in var:
            e1, e2, unit = f"E{r}-D{r}", f"F{r}-D{r}", "bp"
        else:
            e1, e2, unit = f"E{r}/D{r}-1", f"F{r}/D{r}-1", "%"
        # blank (missing) levels must not be read as zero
        f1 = f'=IF(COUNT(D{r},E{r})=2,{e1},"")'
        f2 = f'=IF(COUNT(D{r},F{r})=2,{e2},"")'
        for col, f in [(7, f1), (8, f2)]:
            c = ws.cell(row=r, column=col, value=f)
            c.font = BODY
            c.number_format = "0.0%" if unit == "%" else "0.0"
        ws.cell(row=r, column=9, value=unit).font = BODY
        ws.cell(row=r, column=1).number_format = "yyyy-mm-dd"
    ws.column_dimensions["B"].width = 50


# --------------------------------------------------------------------------- narrative sheets
def summary_sheet() -> None:
    ws = wb.create_sheet("Executive_Summary", 0)
    ws["A1"] = "Gas -> Inflation Expectations -> ECB Pricing -> Bund Curve: Results Recap"
    ws["A1"].font = TITLE
    ws["A2"] = ("Sample: daily Oct 2014 - 9 Oct 2026 (€STR OIS from Oct 2019, 1Y ILS from Apr 2022). Iran-war regime from "
                "2 Mar 2026 (143 trading days / 32 weeks). Green numbers link live to the result sheets.")
    ws["A2"].font = NOTE
    T, W = "Transmission_Daily", "Transmission_Weekly"
    td, tw = FRAMES[T], FRAMES[W]
    war, cris, norm = "Iran war (Mar 2026+)", "2021-22 gas crisis", "2023-Feb26 normalisation"
    rows = [
        # (link, hypothesis, metric, sheet, mask, col, p_sheet, p_col, verdict, meaning)
        ("1. Gas -> realised inflation", "HICP y/y per +10% TTF, 9 months later (pp)",
         "HICP_Passthrough", (FRAMES["HICP_Passthrough"].model.str.startswith("B")) & (FRAMES["HICP_Passthrough"].lag_months == 9),
         "cum_gas_pp_per_10pct", "gas_p", "SUPPORTED, WITH A LAG",
         "Gas reaches consumer prices through retail gas/electricity tariffs that reset slowly; the effect is ~0 on impact "
         "and builds to ~0.2pp of y/y HICP per +10% after 9-12 months. Oil, by contrast, passes through immediately. Gas is "
         "therefore a LEADING indicator of headline inflation at exactly the 1-year horizon the ECB targets."),
        ("1b. ... vs oil (benchmark)", "HICP y/y per +10% Brent, same month (pp)",
         "HICP_Passthrough", (FRAMES["HICP_Passthrough"].model.str.startswith("B")) & (FRAMES["HICP_Passthrough"].lag_months == 0),
         "cum_oil_pp_per_10pct", "oil_p", "CONTEXT",
         "Oil is the faster, larger inflation driver; any 'gas' result must control for Brent (all regressions do)."),
        ("2. Gas -> market inflation expectations", "Infl. swap 1Y, bp per +10% TTF, war regime",
         T, (td.dependent == "Infl. swap 1Y") & (td.regime == war), "bp_per_10pct_ttf", "p_ttf", "SUPPORTED, AMPLIFIED IN WAR",
         "Inflation swaps reprice the same day. Sensitivity in the war (~7bp per +10%) is ~4x the 2023-26 level. The 1y1y "
         "forward moves much less (~1.8bp) and reverses within ~20 trading days (local projections, short sample): markets treat most of the gas spike as a "
         "near-term price-level shock rather than persistent inflation."),
        ("2b. ... forward expectations", "Infl. swap 1y1y, bp per +10% TTF, war regime",
         T, (td.dependent == "Infl. swap 1y1y") & (td.regime == war), "bp_per_10pct_ttf", "p_ttf", "SUPPORTED (WEAKER)",
         "The forward the ECB cares about responds, but modestly, and Brent explains more of it than gas."),
        ("3. Expectations -> ECB rate pricing", "€STR OIS 1y1y, bp per +10% TTF, war regime (weekly)",
         W, (tw.dependent == "€STR OIS 1y1y") & (tw.regime == war), "bp_per_10pct_ttf", "p_ttf", "SUPPORTED",
         "Gas moves the expected policy path in the war (~2.3bp daily, ~4.8bp weekly per +10%). About a third of the daily "
         "effect runs through inflation swaps; the rest is priced directly - the market reads energy headlines straight into "
         "the ECB reaction function. A break test dates the shift to early March 2026 without being told the war date."),
        ("3b. Break in policy sensitivity", "sup-Wald stat, €STR OIS 2Y (5% cv = 8.68)",
         "Structural_Breaks", FRAMES["Structural_Breaks"].dependent == "€STR OIS 2Y", "sup_wald_stat", "robust_wald_p",
         "CONFIRMED (1%)", "The estimated break date is in the first week of the war (see Structural_Breaks)."),
        ("4. Policy pricing -> Bund front end", "Bund 2Y, bp per +10% TTF, war regime (daily; Holm-corrected p within the rates family)",
         T, (td.dependent == "Bund 2Y") & (td.regime == war), "bp_per_10pct_ttf", "p_ttf_holm_family", "SUPPORTED",
         "A +10% gas day lifts the Schatz ~4bp in the war (raw p = 0.003; weekly: ~3.6bp, p = 0.001). After Holm correction "
         "it stays significant within the 13 rate series (p = 0.036) and is borderline across all 20 series tested "
         "(p = 0.057). Weekly mediation: OIS 1y1y passes ~0.7bp per bp into the 2Y Bund."),
        ("4b. Curve shape", "Slope 2s30s, bp per +10% TTF, war regime (daily)",
         T, (td.dependent == "Slope 2s30s") & (td.regime == war), "bp_per_10pct_ttf", "p_ttf", "SUPPORTED: BEAR FLATTENING",
         "The 30Y barely moves (~0.2bp) while the front end rises: gas shocks are priced as POLICY shocks, not as term-premium "
         "or long-run inflation shocks. Confirmed by the PCA slope factor and by weekly 5s30s (-2.6bp, p < 0.001)."),
        ("4c. Which flattener", "Slope 2s30s, weekly, bp per +10% TTF, war",
         "Curve_Slopes", FRAMES["Curve_Slopes"].slope == "2s30s", "war weekly: bp per +10% TTF", "war weekly: p",
         "2s30s LARGEST; 2s10s CLEANEST",
         "All six slopes flatten on gas in the war. 2s30s is the largest and is significant daily and weekly, but Brent "
         "flattens it too. 2s10s (-2.8bp daily) and 2s5s (-2.1bp daily) respond to gas but not to oil, so they are the "
         "purest gas expressions; 2s5s fades within the week. Before 2021 gas slightly STEEPENED the curve (front end "
         "pinned at the lower bound)."),
        ("5. Regime dependence", "Bund 2Y, bp per +10% TTF, 2021-22 gas crisis (daily)",
         T, (td.dependent == "Bund 2Y") & (td.regime == cris), "bp_per_10pct_ttf", "p_ttf", "STRONGLY SUPPORTED",
         "In 2021-22 higher gas meant LOWER Bund and OIS rates: markets priced recession risk and an ECB that called inflation "
         "transitory. In 2026 the ECB is seen as reactive (it hiked in June and September), so gas is a hawkish signal. The "
         "sign of the gas-rates beta is set by the perceived central-bank reaction function, not by gas itself."),
        ("5b. War vs base difference test", "Bund 2Y: war minus base effect, bp per +10%",
         "Regime_Tests", (FRAMES["Regime_Tests"].dependent == "Bund 2Y") & (FRAMES["Regime_Tests"].regime_vs_base == war),
         "difference_bp_per_10pct", "wald_p", "SIGNIFICANT", "Formal test that the war-regime sensitivity differs from normal times."),
        ("6. Predictivity of curve moves", "OOS R2, next-week Bund 2Y, gas model vs no-change (war)",
         "OOS_Forecast", (FRAMES["OOS_Forecast"].target == "Bund 2Y (next week)") & (FRAMES["OOS_Forecast"].evaluation == "Iran war only")
         & (FRAMES["OOS_Forecast"].model == "gas only") & (FRAMES["OOS_Forecast"].benchmark == "no change"),
         "oos_r2", "clark_west_p", "NOT SUPPORTED (market frequency)",
         "Gas EXPLAINS same-day curve moves but does not FORECAST later ones: local projections show no drift after day 1-2, "
         "Granger tests (weekly) find no lead from TTF to Bunds, and OOS forecasts do not beat a random walk. Only next-week "
         "€STR OIS shows a fragile gain in the war (31 forecasts). The predictive channel is macro: gas -> HICP 6-12 months out."),
        ("7. Oil vs gas", "Brent effect on €STR OIS 1y1y, bp per +10%, war (daily)",
         T, (td.dependent == "€STR OIS 1y1y") & (td.regime == war), "bp_per_10pct_brent", "p_brent", "CAVEAT",
         "Brent's effect on inflation and policy pricing is 3-5x gas's. In the Hormuz shock oil and gas co-move (corr ~0.67), "
         "so their separate contributions are only partly identified (VAR FEVD: TTF share of OIS 1y1y variance 2% if Brent is "
         "ordered first, 40% if TTF is first). The gas effect on the Bund front end survives the Brent control."),
    ]
    hdr = ["Link in the chain", "Metric", "Estimate", "p-value", "Verdict", "Economic meaning"]
    r0 = 4
    for j, h in enumerate(hdr, 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font, c.fill, c.alignment = H_FONT, H_FILL, Alignment(wrap_text=True, vertical="center")
    for i, (link, metric, sh, mask, col, pcol, verdict, meaning) in enumerate(rows, r0 + 1):
        ws.cell(row=i, column=1, value=link).font = BOLD
        ws.cell(row=i, column=2, value=metric).font = BODY
        c = ws.cell(row=i, column=3, value=ref(sh, mask, col))
        c.font, c.number_format = LINK, "0.00"
        c = ws.cell(row=i, column=4, value=ref(sh, mask, pcol))
        c.font, c.number_format = LINK, "0.000"
        ws.cell(row=i, column=5, value=verdict).font = BOLD
        ws.cell(row=i, column=6, value=meaning).font = BODY
        for j in range(1, 7):
            ws.cell(row=i, column=j).alignment = WRAP
            ws.cell(row=i, column=j).border = THIN
        ws.row_dimensions[i].height = 92
    ws.conditional_formatting.add(f"D{r0 + 1}:D{r0 + len(rows)}",
                                  CellIsRule(operator="lessThan", formula=["0.05"], fill=PatternFill("solid", fgColor=SIG)))
    for L, w in zip("ABCDEF", [30, 38, 11, 10, 24, 95]):
        ws.column_dimensions[L].width = w
    r = r0 + len(rows) + 2
    blocks = [
        ("Bottom line", [
            "The thesis holds in the current regime: since 2 March 2026, gas has been a significant driver of euro inflation "
            "swaps, of the expected ECB path and of the Bund front end, and it bear-flattens the curve.",
            "It is conditional, not structural: before 2021 the link was ~0 and in 2021-22 it had the opposite sign. Any trade "
            "or model built on it must monitor the rolling sensitivity (Rolling_Betas / chart f3).",
            "Gas explains curve moves contemporaneously; it does not give a tradable lead at daily/weekly frequency. Its genuine "
            "predictive content is for realised inflation 6-12 months ahead, i.e. for the data the ECB will react to.",
        ]),
        ("Methodological safeguards applied", [
            "Model in changes: all levels are I(1) (ADF/KPSS/Zivot-Andrews), no robust cointegration.",
            "Newey-West HAC inference everywhere (fat tails, heteroskedasticity, ARCH and serial correlation detected).",
            "Front-month roll days removed from TTF/Brent/NBP/HH returns (open-interest based detection).",
            "Bundesbank yields are observed ~half a day before LSEG closes: daily Bund regressions include previous-day market "
            "moves; full-chain mediation/VAR with Bunds run weekly; predictive tests skip the overlapping day.",
            "Brent and EURUSD controls in every regression; Holm correction for multiple testing; block-bootstrap CIs.",
        ]),
        ("Limits", [
            "One war episode: 143 daily / 32 weekly observations. Estimates are precise enough to reject zero but not to pin "
            "down magnitudes tightly.",
            "Co-movement conditional on controls is not full causal identification: geopolitical news moves gas, oil and rates "
            "together. A cleaner design would instrument gas with LNG-specific outage news or use intraday event windows.",
            "RESET flags some non-linearity (large gas days may matter disproportionately) - a threshold model is a natural extension.",
        ]),
    ]
    for head, items in blocks:
        ws.cell(row=r, column=1, value=head).font = Font(name=FONT, size=11, bold=True, color=NAVY)
        r += 1
        for it in items:
            ws.cell(row=r, column=1, value="•").font = BODY
            c = ws.cell(row=r, column=2, value=it)
            c.font, c.alignment = BODY, WRAP
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
            ws.row_dimensions[r].height = 30
            r += 1
        r += 1
    ws.freeze_panes = "A5"


def intuition_sheet() -> None:
    ws = wb.create_sheet("Economic_Intuition", 1)
    start = sheet_header(ws, "Economic intuition: the transmission chain and how each link is tested",
                         ["Each row is one link of the thesis: mechanism, expected sign, what would falsify it, the "
                          "benchmark used, the tests, and the outcome."])
    rows = [
        ["1", "Gas price -> realised inflation",
         "Gas sets the marginal price of electricity (merit order) and feeds household gas tariffs, fertiliser, chemicals "
         "and industrial costs. Retail tariffs are regulated or contracted and reset with a lag of months.",
         "+ (lagged)", "No relation between TTF and HICP at any lag once oil is controlled for.",
         "TTF (EUR), HICP y/y and m/m, Brent (EUR)", "Distributed-lag regression (13 lags), Granger, y/y lag scan",
         "Supported: ~0 on impact, +0.2pp y/y per +10% after 9-12 months (p <= 0.02). Oil passes through immediately."],
        ["2", "Gas -> market inflation expectations",
         "Forward-looking investors price the future HICP prints today via inflation-linked swaps (HICPxT). The 1Y spot swap "
         "captures the near-term price level; the 1y1y forward captures persistence into the ECB's horizon.",
         "+", "Inflation swaps do not react to gas once Brent is controlled for.",
         "ILS 1Y (EUCPIZ1Y=TWEB), ILS 1y1y (EUIL1YF1Y=R)", "Regime regressions (HAC), interaction Wald, LP, break tests",
         "Supported. ILS 1Y ~+7bp per +10% TTF in the war (p < 0.001); 1y1y ~+1.8bp; Brent dominates both."],
        ["3", "Inflation expectations -> ECB rate expectations",
         "A central bank with a credible inflation-targeting reaction function raises the expected policy path when expected "
         "inflation rises, especially if second-round effects or de-anchoring are feared.",
         "+ (only if the ECB is seen as reactive)", "OIS does not move with gas/ILS, or moves down (growth fears dominate).",
         "€STR OIS 1M/3M/1Y/2Y, OIS 1y1y forward (derived)", "Mediation with block bootstrap, VAR/Granger, sup-Wald break",
         "Supported in the war: OIS 1y1y +2.3bp (daily) / +4.8bp (weekly) per +10% TTF; ~1/3 via ILS. Negative in 2021-22."],
        ["4", "ECB expectations -> Bund curve",
         "The 2Y Schatz tracks the expected policy path (plus a scarcity premium); long maturities add term premium and "
         "long-run inflation risk. A policy shock lifts the front end more than the long end = bear flattening.",
         "2Y +, slope -", "Long end reacts as much as the front (term-premium shock) or curve steepens.",
         "Bundesbank 2Y-30Y yields, slopes, PCA factors, Bund-OIS spread, 5Y swap spread", "Transmission regressions, PCA, weekly mediation",
         "Supported: 2Y ~+4bp, 30Y ~0, 2s30s ~-3.9bp per +10% TTF; PCA slope factor significant."],
        ["5", "Regime dependence",
         "The SIGN of the gas-rates link depends on how markets expect the ECB to react. If gas is read as a growth shock and "
         "the ECB as patient (2022 'transitory'), yields fall; if the ECB is seen as reactive (2026 hikes), yields rise.",
         "Varies by regime", "Constant beta across regimes.", "All of the above", "Regime splits, interaction tests, Chow/sup-Wald, rolling betas",
         "Strongly supported: negative in 2021-22, ~0 in normal times, positive in the war."],
        ["6", "Predictivity",
         "Efficient markets should price public gas information immediately, so a tradable lead-lag is not expected. "
         "Predictivity is plausible only via slow-moving fundamentals (inflation prints) or limited attention.",
         "None at market frequency; + for macro data", "Gas moves today predict curve moves next days/weeks.",
         "Same series", "Local projections, Granger, out-of-sample Clark-West tests",
         "Not supported for curve moves (no drift after day 1-2, no OOS gain for Bunds); supported for HICP 6-12m ahead."],
        ["C", "Confounders",
         "Oil (Hormuz hits both), EUR exchange rate (energy priced in USD), global risk-off (flight to Bund quality), German "
         "fiscal supply (long end).", "-", "-", "Brent, EURUSD, 5Y swap spread, Bund-OIS spread",
         "Controls in every regression, VAR orderings, swap-spread regressions",
         "Gas effect on front end survives Brent control; oil and gas effects are only partly separable in the war."],
    ]
    df = pd.DataFrame(rows, columns=["#", "Link", "Mechanism (why)", "Expected sign", "What would falsify it",
                                     "Benchmarks used", "Tests", "Outcome"])
    write_df(ws, df, start, widths={"#": 4, "Link": 26, "Mechanism (why)": 60, "Expected sign": 14,
                                     "What would falsify it": 34, "Benchmarks used": 30, "Tests": 30, "Outcome": 48})
    for r in range(start + 1, start + 1 + len(df)):
        ws.row_dimensions[r].height = 105
        for j in range(1, 9):
            ws.cell(row=r, column=j).alignment = WRAP
    r = start + len(df) + 2
    ws.cell(row=r, column=1, value="Chain diagram").font = Font(name=FONT, size=11, bold=True, color=NAVY)
    for k, line in enumerate([
        "Hormuz / Qatar LNG supply shock",
        "   -> TTF gas price (and Brent)",
        "      -> realised HICP (energy, with a 6-12 month lag)       [link 1]",
        "      -> inflation swaps 1Y / 1y1y (priced immediately)      [link 2]",
        "         -> €STR OIS 1y1y = expected ECB path                 [link 3]",
        "            -> Bund 2Y up more than Bund 30Y = bear flattening [link 4]",
    ], 1):
        ws.cell(row=r + k, column=2, value=line).font = Font(name="Courier New", size=10)


def benchmarks_sheet() -> None:
    ws = wb.create_sheet("Benchmarks", 2)
    start = sheet_header(ws, "Benchmarks used, why these, and what was rejected",
                         ["Clean variable names as used in data/clean/*.csv. Full definitions: Data_Dictionary sheet."])
    rows = [
        ["Gas", "gas_ttf_m1_eurmwh", "ICE Endex Dutch TTF front-month (LSEG TFMBMc1)", "EUR/MWh", "Oct 2014 -",
         "THE euro-area wholesale gas benchmark; in EUR, so no FX noise; most liquid European gas contract; front month is "
         "what headlines and retail tariff indexation track.",
         "NBP (UK hub, GBp/therm - kept for robustness), Henry Hub (US, not a euro-area cost), IMF EU monthly average "
         "(monthly, USD, revised).", "Continuation contract: roll days removed (OI jump)."],
        ["Oil (control)", "oil_brent_m1_usdbbl", "ICE Brent front-month (LSEG LCOc1)", "USD/bbl", "Oct 2014 -",
         "Main confounder: Hormuz hits oil and gas together, and inflation swaps are historically oil-driven.",
         "WTI (US benchmark, less relevant for Europe).", "Roll days removed."],
        ["FX (control)", "fx_eurusd", "EUR/USD spot (LSEG EUR=)", "USD per EUR", "Oct 2014 -",
         "Energy is priced in USD; EUR moves change the euro cost of oil and risk sentiment.", "-", "Mid of bid/ask."],
        ["Inflation expectations (near)", "ils_eur_1y", "EUR 1Y zero-coupon inflation swap, HICPxT (LSEG/Tradeweb EUCPIZ1Y=TWEB)",
         "%", "Apr 2022 -", "Captures expected inflation over the next year: where gas pass-through shows up first.",
         "HICP flash surveys (monthly, not market-based).", "Short history (from 2022)."],
        ["Inflation expectations (forward)", "ils_eur_1y1y", "EUR inflation swap 1y forward 1y (LSEG EUIL1YF1Y=R)", "%",
         "Jan 2016 -", "Inflation 1-2 years ahead = the ECB's policy horizon; strips out the mechanical near-term carry.",
         "5y5y (long-run anchoring - not available, suggested addition).", "File was mislabelled '1Y'; renamed."],
        ["Rate-hike expectations", "ois_estr_1y1y / ois_estr_2y", "€STR OIS 1M, 3M, 1Y, 2Y (LSEG EUREST*=); 1y1y derived",
         "%", "Sep 2019 -", "OIS is the cleanest market measure of the expected ECB path: no credit, scarcity or "
         "collateral premium (unlike Bunds). 1y1y = expected rate one year ahead.",
         "Euribor futures (contain bank credit premium), ECB-dated OIS (not available).", "Mid of bid/ask."],
        ["Bund curve", "bund_2y ... bund_30y", "Bundesbank daily yields of current federal securities (BBSSY)", "%",
         "2Y/5Y/10Y/30Y from 2014; 7Y/15Y from 2020; 20Y from Jan 2026",
         "Official, free, long history. 2Y Schatz = policy-sensitive end; 30Y = term premium / long-run end.",
         "Zero-coupon fitted curve (would remove on-the-run roll effects - suggested robustness).",
         "Snapshot taken ~half a day before LSEG closes: timing-adjusted."],
        ["Curve shape", "slope_2s10s, slope_2s30s, slope_5s30s, fly_2s5s10s, PCA", "Derived from Bund yields", "bp", "-",
         "Tests the 'specific moves' part of the thesis: policy shocks flatten, term-premium shocks steepen.", "-", "-"],
        ["Safety / scarcity premia", "swapsprd_de5y_bp, sprd_bund2y_ois2y", "5Y IRS - Bund (LSEG DE5IRS5Y=RR); Bund 2Y - OIS 2Y",
         "bp", "2014 - / 2019 -", "Separates genuine policy repricing from flight-to-quality into Bunds.", "-", "-"],
        ["Realised inflation", "hicp_ea_idx / _yoy / _mom", "Euro-area HICP index (2025=100)", "index / %", "Dec 1999 -",
         "First link of the chain at its natural (monthly) frequency.", "HICP energy sub-index (suggested addition).",
         "Base-splice break removed; NSA - month dummies used."],
        ["Robustness only", "gas_nbp_m1_*, gas_hh_m1_usdmmbtu, gas_eu_imf_usdmmbtu_mavg", "NBP, Henry Hub, IMF EU gas",
         "various", "various", "Cross-checks; HH is the US leg of the LNG arbitrage.", "-", "Not used in headline tests."],
    ]
    df = pd.DataFrame(rows, columns=["Role in the chain", "Clean name", "Instrument / source", "Unit", "History",
                                     "Why this benchmark", "Alternatives considered", "Treatment / caveat"])
    write_df(ws, df, start, widths={"Role in the chain": 24, "Clean name": 26, "Instrument / source": 40, "Unit": 10,
                                     "History": 18, "Why this benchmark": 60, "Alternatives considered": 40,
                                     "Treatment / caveat": 34})
    for r in range(start + 1, start + 1 + len(df)):
        ws.row_dimensions[r].height = 75
        for j in range(1, 9):
            ws.cell(row=r, column=j).alignment = WRAP


def catalog_sheet() -> None:
    ws = wb.create_sheet("Test_Catalog", 3)
    start = sheet_header(ws, "Catalog of econometric tests: what each answers, result, and economic meaning",
                         ["Ordered from data properties to identification to predictivity. 'Sheet' points to the full table."])
    rows = [
        ["Data properties", "Jarque-Bera", "Are daily changes normal?", "Normality",
         "Rejected for all series (excess kurtosis up to 15).", "Fat tails -> rely on robust (HAC) inference and bootstrap CIs.", "Descriptive"],
        ["Data properties", "Ljung-Box Q(10)", "Are changes autocorrelated?", "No autocorrelation",
         "Rejected in the full sample for 2Y Bund, TTF, ILS; not in the war regime.", "Mild serial dependence -> Newey-West SEs.", "Descriptive"],
        ["Data properties", "ARCH-LM(5)", "Volatility clustering?", "No ARCH", "Strongly rejected in full sample.",
         "Volatility regimes -> HAC SEs; regime split.", "Descriptive"],
        ["Stationarity", "ADF / KPSS / Zivot-Andrews", "Levels or changes?", "Unit root (ADF, ZA) / stationary (KPSS)",
         "All levels I(1); all changes I(0) (OIS 2Y borderline on ADF only).", "Regressions in levels would be spurious -> model changes.", "Unit_Roots"],
        ["Long-run relation", "Engle-Granger, Johansen", "Do gas and yields share a long-run equilibrium?", "No cointegration",
         "No pairwise cointegration; Johansen rank > 0 is likely break-driven (2022).", "No stable levels relationship: gas is a shock driver, not an anchor for yields.", "Cointegration"],
        ["Transmission", "OLS with Newey-West HAC, by regime", "How much do rates move per gas move?", "beta = 0",
         "War: ILS 1Y, OIS, Bund 2Y, 2s30s significant; 2021-22 negative; normal times ~0.", "Core result: gas is a front-end policy driver in the war.", "Transmission_Daily / _Weekly"],
        ["Multiple testing", "Holm correction", "Are results artefacts of testing many series?", "All betas = 0",
         "War regime: ILS 1Y survives across all 20 series (p = 0.005); Bund 2Y survives within the rates family "
         "(p = 0.036), borderline across all 20 (p = 0.057); 2s5s survives within the curve-shape family (p = 0.045).",
         "The headline is not a data-mining artefact, but individual slope results are suggestive.", "Transmission_Daily"],
        ["Regime", "Interaction Wald test", "Is the war beta different from normal?", "Equal betas",
         "Rejected for Bund 2Y, ILS 1Y, 2s10s, 2s30s.", "The link is regime-specific.", "Regime_Tests"],
        ["Regime", "Chow, HAC Wald, Quandt-Andrews sup-Wald", "Is there a break, and when?", "No break",
         "OIS 1y1y / OIS 2Y: break dated 5-11 Mar 2026 (1%). Slopes: break in Feb-Mar 2025 (German fiscal shock).",
         "The data locate the regime change at the war onset on their own.", "Structural_Breaks"],
        ["Regime", "CUSUM", "Are full-sample coefficients stable?", "Stable", "Rejected for Bund levels in the full sample.",
         "A single full-sample beta is misleading.", "Diagnostics"],
        ["Residuals", "Durbin-Watson, Breusch-Godfrey, Ljung-Box", "Serial correlation in residuals?", "None",
         "Present in the full sample; in the war regime only for OIS 1y1y and Bund 10Y.", "Handled by HAC SEs.", "Diagnostics"],
        ["Residuals", "Breusch-Pagan, White, ARCH-LM", "Heteroskedasticity?", "Homoskedastic", "Present.", "Handled by HAC SEs.", "Diagnostics"],
        ["Residuals", "Ramsey RESET", "Functional form?", "Linear is adequate", "Rejected in most full-sample models and for Bund 2Y and 2s10s in the war regime.",
         "Possible non-linearity: big gas days matter more -> extension: threshold model.", "Diagnostics"],
        ["Residuals", "VIF", "Gas vs oil collinearity?", "-", "VIF < 2.5 (corr TTF-Brent ~0.68 in the war).",
         "Collinearity is moderate: separate effects identifiable but noisier.", "Diagnostics"],
        ["Timing", "Lead/lag vs OIS; Dimson", "Are data sources synchronous?", "-",
         "Bundesbank yields load ~50/50 on previous- and same-day OIS moves.",
         "Bundesbank snapshot ~half a day early -> timing-adjusted specifications (prevents spurious 'predictivity').", "Timing"],
        ["Channel", "Mediation (block bootstrap)", "Does gas move rates via inflation expectations?", "Indirect effect = 0",
         "Daily war: ~32% of TTF->OIS 1y1y via ILS 1y1y; weekly war: total indirect via ILS/OIS to Bund 2Y significant.",
         "The inflation-expectations channel is real but the market also prices the ECB reaction directly.", "Mediation"],
        ["Channel / prediction", "VAR, Granger, IRF, FEVD", "Dynamic interactions and lead-lag", "No Granger causality",
         "Daily full sample: TTF Granger-causes OIS 1y1y/2Y (p < 0.01) but the LP shows it is a small next-day give-back (~-0.8bp per +10%), not momentum. No TTF lead in the war or at weekly frequency. FEVD bounds TTF share of OIS 1y1y at 2-40% (ordering).", "Contemporaneous, not lagged, transmission.", "Granger / VAR_FEVD"],
        ["Prediction", "Local projections (Jorda)", "Does gas forecast LATER curve moves?", "No predictive effect",
         "No significant drift after day 0-1 for any rate; ILS 1y1y reverses within ~20 days in the war (suggestive).",
         "Markets absorb gas news within 1-2 days; no tradable momentum.", "Local_Projections"],
        ["Prediction", "Out-of-sample, Clark-West", "Does a gas model beat a random walk next week?", "Equal accuracy",
         "No for Bunds and ILS; fragile yes for OIS in the war (31 forecasts).", "No robust forecasting edge.", "OOS_Forecast"],
        ["Macro link", "Distributed lag + y/y lag scan", "Does gas predict realised inflation?", "No effect at any lag",
         "Joint lags significant; peak at 9-12 months.", "Gas is a leading indicator of the inflation prints the ECB reacts to.", "HICP_Passthrough"],
        ["Curve", "Slope comparison (2s5s, 2s10s, 5s10s, 10s30s, 2s30s, 5s30s)", "Which flattener carries gas?", "beta = 0",
         "All flatten in the war; 2s30s largest and robust; 2s10s / 2s5s gas-specific (no oil effect); 2s5s fades weekly.",
         "Choose 2s10s for a pure gas view, 2s30s for size (but it carries oil risk).", "Curve_Slopes"],
        ["Curve", "PCA (level / slope / curvature)", "Which curve factor does gas move?", "-",
         "War: slope factor significantly negative (flattening); level positive but less precise.", "Policy shock signature.", "Curve_PCA"],
        ["Events", "Event study", "What happened on key dates?", "-", "War onset (2-day): Bund 2Y +16bp vs 30Y +6bp, 2s10s -5bp. Ceasefire: OIS 1y1y -23bp, 2s10s +7bp (reversal).",
         "Narrative check of the regressions.", "Event_Study"],
    ]
    df = pd.DataFrame(rows, columns=["Category", "Test", "Question", "Null hypothesis (H0)", "Result",
                                     "Economic meaning / consequence", "Sheet"])
    write_df(ws, df, start, widths={"Category": 18, "Test": 30, "Question": 34, "Null hypothesis (H0)": 24,
                                     "Result": 52, "Economic meaning / consequence": 52, "Sheet": 22})
    for r in range(start + 1, start + 1 + len(df)):
        ws.row_dimensions[r].height = 48
        for j in range(1, 8):
            ws.cell(row=r, column=j).alignment = WRAP


def data_sheets() -> None:
    for name, csv, title in [("Data_Quality", "data_quality_log", "Data quality log: issues found and treatment"),
                             ("Data_Dictionary", "data_dictionary", "Data dictionary (clean variable names)"),
                             ("Coverage", "coverage", "Series coverage")]:
        df = pd.read_csv(CLEAN / f"{csv}.csv")
        if name == "Coverage":
            df = df.rename(columns={"Unnamed: 0": "series"})
        ws = wb.create_sheet(name)
        start = sheet_header(ws, title, [f"Source: data/clean/{csv}.csv (built by src/clean.py)."])
        write_df(ws, df, start, widths={"issue": 70, "treatment": 70, "description": 70})
        for r in range(start + 1, start + 1 + len(df)):
            for j in range(1, len(df.columns) + 1):
                ws.cell(row=r, column=j).alignment = WRAP


def charts_sheet() -> None:
    ws = wb.create_sheet("Charts", 4)
    ws["A1"] = "Charts"
    ws["A1"].font = TITLE
    ws["A2"] = "Rendered by src/figures.py; PNGs in output/figures/."
    ws["A2"].font = NOTE
    row = 4
    for f in ["f1_war_timeline", "f2_regime_betas", "f6_flatteners", "f7_tenor_moves_war", "f3_rolling_beta",
              "f4_local_projections", "f5_hicp_lag_scan"]:
        img = Image(str(FIGURES / f"{f}.png"))
        scale = 820 / img.width
        img.width, img.height = int(img.width * scale), int(img.height * scale)
        ws.add_image(img, f"A{row}")
        row += int(img.height / 20) + 3


def readme_sheet() -> None:
    ws = wb.create_sheet("README", 0)
    ws["A1"] = "How to read this workbook"
    ws["A1"].font = TITLE
    lines = [
        ("Purpose", "Test whether European gas prices drive / predict moves on the German yield curve through inflation "
                    "expectations and ECB rate expectations, with focus on the Iran-war regime (from 2 Mar 2026)."),
        ("Start here", "Executive_Summary (results & economic meaning) -> Economic_Intuition (the chain) -> Benchmarks "
                       "(why these instruments) -> Test_Catalog (every test, one line each) -> Charts."),
        ("Result sheets", "One sheet per test family, exported from Python. Green fill = p < 0.05 (in Diagnostics: a problem detected)."),
        ("Units", "Rates: bp. Prices: % log change. 'bp_per_10pct_ttf' = bp move for a +10% TTF move."),
        ("Colour code", "Green text = live formula linking to another sheet. Black = values/text. Formulas in Event_Study "
                        "compute changes from levels."),
        ("Reproduce", ".venv/bin/python src/clean.py && .venv/bin/python src/econometrics.py && "
                      ".venv/bin/python src/figures.py && .venv/bin/python src/build_workbook.py"),
        ("Data", "Raw exports archived under canonical names in data/raw/; clean panels in data/clean/."),
        ("Prepared", "11 Oct 2026, data to 9 Oct 2026 close."),
    ]
    for i, (k, v) in enumerate(lines, 3):
        ws.cell(row=i, column=1, value=k).font = BOLD
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.alignment = BODY, WRAP
        ws.row_dimensions[i].height = 34
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 120


def main() -> None:
    build_results()
    summary_sheet()
    intuition_sheet()
    benchmarks_sheet()
    catalog_sheet()
    charts_sheet()
    data_sheets()
    readme_sheet()
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # no cached values from openpyxl: Excel computes on open
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
