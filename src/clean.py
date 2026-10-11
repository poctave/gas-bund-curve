"""Load raw exports, archive them under canonical names, clean, and build panels.

Outputs (data/clean/):
  daily_levels.csv     business-day levels, all series, + regime / roll flags
  daily_changes.csv    bp changes (rates) and 100*log changes (prices)
  weekly_changes.csv   Friday-to-Friday changes
  monthly_panel.csv    HICP, IMF gas, month-average and month-end market data
  data_dictionary.csv  one row per variable
  data_quality_log.csv issues found and their treatment
"""
from __future__ import annotations

import glob
import shutil

import numpy as np
import openpyxl
import pandas as pd

from config import (BUND_CODES, CLEAN, HICP_RAW, HICP_SOURCE, ILS_1Y1Y_RAW, ILS_1Y1Y_SOURCE,
                    IMF_GAS_RAW, IMF_GAS_SOURCE, LSEG_FILES, RAW, ROOT, THERM_PER_MWH, regime_of)

QUALITY_LOG: list[dict] = []


def log_issue(series: str, issue: str, treatment: str, n_affected: int | str = "") -> None:
    QUALITY_LOG.append({"series": series, "issue": issue, "treatment": treatment,
                        "n_affected": n_affected})


# --------------------------------------------------------------------------- raw archive
def archive_raw() -> None:
    """Copy every source file into data/raw/ under a canonical name (originals untouched)."""
    for f in glob.glob(str(ROOT / "Price History_*.xlsx")):
        ws = openpyxl.load_workbook(f, read_only=True, data_only=True).worksheets[0]
        ric = next(ws.iter_rows(min_row=4, max_row=4, values_only=True))[0]
        if ric not in LSEG_FILES:
            continue
        dest = RAW / LSEG_FILES[ric][0]
        # Several exports of the same RIC exist (NBP 2y and 10y): keep the longest.
        if not dest.exists() or _n_rows(f) > _n_rows(dest):
            shutil.copy2(f, dest)
    shutil.copy2(ROOT / ILS_1Y1Y_SOURCE, RAW / ILS_1Y1Y_RAW)
    shutil.copy2(ROOT / IMF_GAS_SOURCE, RAW / IMF_GAS_RAW)
    shutil.copy2(ROOT / HICP_SOURCE, RAW / HICP_RAW)
    for f in glob.glob(str(ROOT / "BBSSY*.csv")):
        tenor = BUND_CODES[f.split(".")[4]]
        shutil.copy2(f, RAW / f"bbk_bund_{tenor}_{f.split('/')[-1]}")


def _n_rows(path) -> int:
    return openpyxl.load_workbook(path, read_only=True).worksheets[0].max_row


# --------------------------------------------------------------------------- parsers
def read_lseg(path) -> pd.DataFrame:
    """Parse an LSEG 'Price History' export. Returns DataFrame indexed by date.

    Futures exports carry Close/OI; swap/FX exports carry Bid/Ask -> mid.
    """
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)
                .worksheets[0].iter_rows(values_only=True))
    i = next(k for k, r in enumerate(rows) if r[0] in ("Exchange Date", "Date")
             and (r[0] == "Exchange Date" or r[1] == "Exchange Date"))
    header = [h for h in rows[i] if h is not None]
    df = pd.DataFrame([r[:len(header)] for r in rows[i + 1:]], columns=header)
    date_col = "Exchange Date"
    df = df[pd.to_datetime(df[date_col], errors="coerce").notna()].copy()
    df.index = pd.to_datetime(df[date_col]).dt.normalize()
    df = df.sort_index()
    out = pd.DataFrame(index=df.index)
    if "Close" in df:
        out["px"] = pd.to_numeric(df["Close"], errors="coerce")
        out["oi"] = pd.to_numeric(df.get("OI"), errors="coerce")
    else:
        bid = pd.to_numeric(df["Bid"], errors="coerce")
        ask = pd.to_numeric(df["Ask"], errors="coerce")
        out["px"] = (bid + ask) / 2
        out["bid_ask"] = ask - bid
    return out[~out.index.duplicated(keep="last")]


def read_bundesbank(path) -> pd.DataFrame:
    d = pd.read_csv(path, skiprows=9, header=None, names=["date", "y", "flag"], encoding="utf-8-sig")
    d["date"] = pd.to_datetime(d["date"], errors="coerce")
    d["y"] = pd.to_numeric(d["y"], errors="coerce")
    return d.dropna(subset=["date"]).set_index("date")


def detect_rolls(oi: pd.Series, fallback: str, threshold: float = 1.25) -> pd.Series:
    """One front-month roll per month boundary.

    Within a window of the last 4 business days of month M and first 3 of M+1, the roll is the
    day open interest jumps most (expiring contract's OI has run off; the new front month carries
    far more), if that jump exceeds `threshold`. Otherwise fall back to the exchange calendar:
    'last' = last business day of M (ICE Endex TTF, NBP, NYMEX HH),
    'first' = first business day of M+1 (ICE Brent).
    """
    idx = oi.index
    ratio = (oi / oi.shift(1)).fillna(0)
    month = idx.to_period("M")
    out = pd.Series(False, index=idx)
    for m in month.unique()[:-1]:
        cur = np.where(month == m)[0]
        nxt = np.where(month == m + 1)[0]
        if len(cur) == 0 or len(nxt) == 0:
            continue
        win = np.r_[cur[-4:], nxt[:3]]
        best = win[np.argmax(ratio.iloc[win].values)]
        if ratio.iloc[best] > threshold:
            out.iloc[best] = True
        else:
            out.iloc[cur[-1] if fallback == "last" else nxt[0]] = True
    return out


# --------------------------------------------------------------------------- build
def build_daily() -> tuple[pd.DataFrame, pd.DataFrame]:
    # Bund yields
    bunds, flags = {}, {}
    for f in glob.glob(str(ROOT / "BBSSY*.csv")):
        tenor = BUND_CODES[f.split(".")[4]]
        d = read_bundesbank(f)
        s = d["y"].dropna()
        bunds[f"bund_{tenor}"] = s
        fl = d.loc[d["y"].notna(), "flag"].dropna()
        flags[f"bund_{tenor}"] = fl
        if len(fl):
            log_issue(f"bund_{tenor}", f"Bundesbank quality flags: {fl.value_counts().to_dict()}",
                      "Kept (flags are informational); listed in daily_levels *_flag columns", len(fl))
        n_na = int(d["y"].isna().sum())
        log_issue(f"bund_{tenor}", "'.' placeholders for weekends/holidays (calendar-day file)",
                  "Dropped; panel built on Bundesbank business days", n_na)
    bunds = pd.DataFrame(bunds)
    cal = bunds.index[bunds[["bund_2y", "bund_10y", "bund_30y"]].notna().any(axis=1)]
    cal = cal[cal >= "2014-01-01"]
    panel = pd.DataFrame(index=cal)
    panel.index.name = "date"
    for c in ["bund_2y", "bund_5y", "bund_7y", "bund_10y", "bund_15y", "bund_20y", "bund_30y"]:
        panel[c] = bunds[c].reindex(cal)
    for c, fl in flags.items():
        panel[f"{c}_flag"] = fl.reindex(cal)

    rolls = {}
    for ric, (raw_name, var) in LSEG_FILES.items():
        d = read_lseg(RAW / raw_name)
        s = d["px"]
        if var == "swapsprd_de5y_bp" or var.startswith("ois_") or var == "ils_eur_1y":
            pass
        n_missing = int(s.isna().sum())
        if n_missing:
            log_issue(var, "Missing bid or ask / close in export", "Left as NaN (no forward fill)", n_missing)
        extra = d.index.difference(cal)
        if len(extra):
            log_issue(var, "Dates not in Bundesbank business-day calendar (UK/US trading days)",
                      "Dropped from daily panel (calendar = German bond market days)", len(extra))
        panel[var] = s.reindex(cal)
        if "oi" in d and d["oi"].notna().sum() > 100:
            r = detect_rolls(d["oi"], fallback="first" if ric == "LCOc1" else "last")
            rolls[var] = r.reindex(cal, fill_value=False)
            log_issue(var, "Front-month continuation: price jumps at contract roll",
                      "One roll per month: largest open-interest jump (>1.25x) around month-end, "
                      "else exchange-calendar roll day; roll-day returns excluded from regressions "
                      "(flag roll_<var>; *_raw returns kept)", int(r.sum()))
        if "bid_ask" in d:
            ba = d["bid_ask"]
            wide = ba > ba.quantile(0.99) * 3
            if wide.any():
                log_issue(var, "Abnormally wide bid/ask (>3x 99th pct)", "Mid kept; flagged in log", int(wide.sum()))

    ils = pd.read_excel(ROOT / ILS_1Y1Y_SOURCE, header=None, skiprows=2, names=["date", "v"]).dropna()
    ils["date"] = pd.to_datetime(ils["date"])
    panel["ils_eur_1y1y"] = ils.set_index("date")["v"].astype(float).reindex(cal)
    log_issue("ils_eur_1y1y", "File named '1Y inflation swap' is RIC EUIL1YF1Y=R, i.e. the 1y1y FORWARD",
              "Renamed ils_eur_1y1y; spot 1Y added separately from EUCPIZ1Y=TWEB", "")
    log_issue("ils_eur_1y", "Tradeweb 1Y ILS history only starts 2022-04-08", "Used in post-2022 tests only", "")

    # Derived series
    panel["ois_estr_1y1y"] = (((1 + panel.ois_estr_2y / 100) ** 2 / (1 + panel.ois_estr_1y / 100)) - 1) * 100
    panel["gas_nbp_m1_eurmwh"] = panel.gas_nbp_m1_gbppth / 100 * THERM_PER_MWH / panel.fx_eurgbp
    panel["oil_brent_m1_eurbbl"] = panel.oil_brent_m1_usdbbl / panel.fx_eurusd
    panel["gas_ttf_m1_usdmmbtu"] = panel.gas_ttf_m1_eurmwh * panel.fx_eurusd / 3.41214
    panel["slope_2s5s"] = (panel.bund_5y - panel.bund_2y) * 100
    panel["slope_2s10s"] = (panel.bund_10y - panel.bund_2y) * 100
    panel["slope_5s10s"] = (panel.bund_10y - panel.bund_5y) * 100
    panel["slope_10s30s"] = (panel.bund_30y - panel.bund_10y) * 100
    panel["slope_2s30s"] = (panel.bund_30y - panel.bund_2y) * 100
    panel["slope_5s30s"] = (panel.bund_30y - panel.bund_5y) * 100
    panel["fly_2s5s10s"] = (2 * panel.bund_5y - panel.bund_2y - panel.bund_10y) * 100
    panel["sprd_bund2y_ois2y"] = (panel.bund_2y - panel.ois_estr_2y) * 100
    panel["ttf_brent_ratio"] = panel.gas_ttf_m1_usdmmbtu / (panel.oil_brent_m1_usdbbl / 5.8)
    panel["regime"] = regime_of(panel.index).values
    panel["war"] = (panel.index >= "2026-03-02").astype(int)
    for var, r in rolls.items():
        panel[f"roll_{var}"] = r.astype(int)
    return panel, bunds


RATE_VARS = ["bund_2y", "bund_5y", "bund_7y", "bund_10y", "bund_15y", "bund_20y", "bund_30y",
             "ois_estr_1m", "ois_estr_3m", "ois_estr_1y", "ois_estr_2y", "ois_estr_1y1y",
             "ils_eur_1y", "ils_eur_1y1y"]
BP_VARS = ["slope_2s5s", "slope_2s10s", "slope_5s10s", "slope_10s30s", "slope_2s30s", "slope_5s30s", "fly_2s5s10s", "sprd_bund2y_ois2y",
           "swapsprd_de5y_bp"]
PRICE_VARS = ["gas_ttf_m1_eurmwh", "oil_brent_m1_usdbbl", "oil_brent_m1_eurbbl", "gas_nbp_m1_gbppth",
              "gas_nbp_m1_eurmwh", "gas_hh_m1_usdmmbtu", "fx_eurusd", "fx_eurgbp"]


def to_changes(levels: pd.DataFrame) -> pd.DataFrame:
    ch = pd.DataFrame(index=levels.index)
    for c in RATE_VARS:
        ch[f"d_{c}"] = levels[c].diff() * 100  # bp
    for c in BP_VARS:
        ch[f"d_{c}"] = levels[c].diff()  # already bp
    for c in PRICE_VARS:
        ch[f"r_{c}"] = np.log(levels[c]).diff() * 100  # % log return
    return ch


def build_daily_changes(panel: pd.DataFrame) -> pd.DataFrame:
    ch = to_changes(panel)
    # roll-contaminated returns -> NaN in a separate "clean" column; raw kept for robustness
    for var in ["gas_ttf_m1_eurmwh", "oil_brent_m1_usdbbl", "gas_nbp_m1_gbppth", "gas_hh_m1_usdmmbtu"]:
        rc = f"roll_{var}"
        if rc in panel:
            ch[f"r_{var}_raw"] = ch[f"r_{var}"]
            ch.loc[panel[rc] == 1, f"r_{var}"] = np.nan
    ch.loc[panel["roll_oil_brent_m1_usdbbl"] == 1, "r_oil_brent_m1_eurbbl"] = np.nan
    ch.loc[panel["roll_gas_nbp_m1_gbppth"] == 1, "r_gas_nbp_m1_eurmwh"] = np.nan
    ch["regime"] = panel["regime"]
    ch["war"] = panel["war"]
    return ch


def build_weekly_changes(panel: pd.DataFrame, daily: pd.DataFrame) -> pd.DataFrame:
    """Friday-to-Friday changes.

    Rates: change in Friday levels. Prices: sum of the daily roll-cleaned log returns, so contract
    rolls never enter weekly returns. `*_lag1d` columns sum daily returns shifted one business day
    (Friday-to-Thursday window): used when the target is a Bundesbank yield, whose daily snapshot is
    taken earlier than the market closes (see lead/lag tests)."""
    num = panel.select_dtypes("number")
    lv = num.resample("W-FRI").last()
    ch = to_changes(lv)
    rcols = [c for c in daily if c.startswith("r_") and not c.endswith("_raw")]
    ch[rcols] = daily[rcols].resample("W-FRI").sum(min_count=3)
    lagged = daily[rcols].shift(1).resample("W-FRI").sum(min_count=3).add_suffix("_lag1d")
    ch = ch.join(lagged)
    ch["regime"] = regime_of(ch.index).values
    ch["war"] = (ch.index >= "2026-03-06").astype(int)
    return ch


def build_monthly(panel: pd.DataFrame) -> pd.DataFrame:
    h = pd.read_csv(ROOT / HICP_SOURCE)
    h["date"] = pd.to_datetime(h["Date"], dayfirst=True).dt.to_period("M").dt.to_timestamp()
    h = h.set_index("date")["HICP"]
    brk = h.pct_change().idxmin()
    log_issue("hicp_ea_idx", f"Two index bases spliced: fake {h.pct_change().min() * 100:.1f}% m/m at {brk:%Y-%m}",
              "Truncated: only data from the break onward (2025=100 base) kept", int((h.index < brk).sum()))
    h = h[h.index >= brk]
    log_issue("hicp_ea_idx", "Not seasonally adjusted", "Use y/y, or m/m with calendar-month dummies", "")
    m = pd.DataFrame({"hicp_ea_idx": h})
    m["hicp_ea_yoy"] = m.hicp_ea_idx.pct_change(12) * 100
    m["hicp_ea_mom"] = m.hicp_ea_idx.pct_change() * 100
    g = pd.read_csv(ROOT / IMF_GAS_SOURCE, parse_dates=[0], index_col=0).iloc[:, 0]
    m = m.join(g.rename("gas_eu_imf_usdmmbtu_mavg"), how="outer")
    log_issue("gas_eu_imf_usdmmbtu_mavg", "Monthly average, USD, revised by FRED, ends 2026-07",
              "Used only for monthly cross-checks; TTF (EUR) is the primary gas variable", "")
    num = panel.select_dtypes("number").drop(columns=[c for c in panel if c.startswith(("roll_", "war"))])
    mavg = num.resample("MS").mean().add_suffix("_mavg")
    mend = num.resample("MS").last().add_suffix("_mend")
    m = m.join(mavg, how="left").join(mend, how="left")
    m.index.name = "month"
    return m


DICT_ROWS = [
    # name, description, unit, source / RIC, frequency
    ("bund_{t}", "Yield of the current (on-the-run) German federal security, {t} maturity bucket", "%", "Bundesbank BBSSY.D.REN.EUR.A6xx", "daily"),
    ("ois_estr_1m/3m/1y/2y", "EUR €STR overnight index swap par rate", "%", "LSEG EUREST1M=/3M=/1Y=/2Y= (mid)", "daily, from 2019-09"),
    ("ois_estr_1y1y", "1-year €STR rate 1 year forward = ((1+OIS2y)^2/(1+OIS1y))-1", "%", "Derived", "daily"),
    ("ils_eur_1y", "EUR zero-coupon inflation swap 1Y (HICPxT)", "%", "LSEG/Tradeweb EUCPIZ1Y=TWEB (mid)", "daily, from 2022-04"),
    ("ils_eur_1y1y", "EUR inflation swap 1y forward 1y (HICPxT)", "%", "LSEG EUIL1YF1Y=R (mid)", "daily, from 2016"),
    ("gas_ttf_m1_eurmwh", "ICE Endex Dutch TTF front-month futures settlement", "EUR/MWh", "LSEG TFMBMc1", "daily"),
    ("gas_ttf_m1_usdmmbtu", "TTF converted at EURUSD; 1 MWh = 3.41214 MMBtu", "USD/MMBtu", "Derived", "daily"),
    ("gas_nbp_m1_gbppth", "ICE UK NBP front-month futures", "GBp/therm", "LSEG NGLNMc1", "daily, from 2016-11"),
    ("gas_nbp_m1_eurmwh", "NBP converted: p/th /100 x 34.12 th/MWh / EURGBP", "EUR/MWh", "Derived", "daily"),
    ("gas_hh_m1_usdmmbtu", "NYMEX Henry Hub front-month futures", "USD/MMBtu", "LSEG NGc1", "daily, from 2024-10"),
    ("oil_brent_m1_usdbbl", "ICE Brent front-month futures", "USD/bbl", "LSEG LCOc1", "daily"),
    ("oil_brent_m1_eurbbl", "Brent in EUR", "EUR/bbl", "Derived", "daily"),
    ("fx_eurusd / fx_eurgbp", "Spot FX mid", "USD per EUR / GBP per EUR", "LSEG EUR= / EURGBP=", "daily"),
    ("swapsprd_de5y_bp", "5Y EUR IRS minus 5Y Bund benchmark yield (swap spread; Bund scarcity/safety premium)", "bp", "LSEG DE5IRS5Y=RR (mid)", "daily"),
    ("slope_2s5s / 2s10s / 5s10s / 10s30s / 2s30s / 5s30s", "Bund curve slopes (long minus short)", "bp", "Derived", "daily"),
    ("fly_2s5s10s", "Butterfly 2*5Y - 2Y - 10Y (curvature)", "bp", "Derived", "daily"),
    ("sprd_bund2y_ois2y", "2Y Bund minus 2Y €STR OIS (Schatz scarcity / collateral premium)", "bp", "Derived", "daily"),
    ("ttf_brent_ratio", "TTF (USD/MMBtu) / Brent per MMBtu (bbl/5.8): gas vs oil relative price", "ratio", "Derived", "daily"),
    ("regime", "R0 pre-2021 / R1 2021-22 gas crisis / R2 2023-Feb26 / R3 Iran war (from 2026-03-02)", "label", "Defined in src/config.py", "daily"),
    ("roll_<var>", "1 on front-month contract roll days (OI jump)", "0/1", "Derived", "daily"),
    ("d_<rate>", "Daily change in rate", "bp", "Derived", "daily"),
    ("r_<price>", "Daily log return x100 (roll days set to NaN; *_raw keeps them)", "%", "Derived", "daily"),
    ("hicp_ea_idx / _yoy / _mom", "Euro-area HICP index (2025=100, NSA), y/y and m/m", "index / %", "eurozone_hicp.csv", "monthly"),
    ("gas_eu_imf_usdmmbtu_mavg", "IMF Global price of natural gas, EU (monthly average)", "USD/MMBtu", "FRED PNGASEUUSDM", "monthly"),
]


def main() -> None:
    archive_raw()
    panel, _ = build_daily()
    daily = build_daily_changes(panel)
    weekly = build_weekly_changes(panel, daily)
    monthly = build_monthly(panel)
    log_issue("stock_data.csv", "Unknown source; weekend rows duplicating Fridays; High<Open rows; "
              "does not match NGc1 closes", "Excluded from all analysis", 24)
    log_issue("Price History_20261010_1208.xlsx", "2-year NBP export is an exact subset of the 10-year export",
              "Excluded (redundant)", 506)
    log_issue("bund_*", "Asynchronous snapshot: daily Bund change loads ~equally on the previous and the "
              "same day's LSEG OIS change (0.53 / 0.49, t>10); corr(Bund2Y, OIS2Y) rises from 0.44 "
              "(1-day) to 0.67 (2-day) to 0.87 (weekly) => Bundesbank yields are observed roughly half a "
              "trading day before LSEG closes / TTF settlement",
              "Daily Bund regressions include same-day and previous-day market moves (Dimson); "
              "total effect = sum. Full-chain mediation and VAR with Bunds run weekly; predictive "
              "tests for Bund targets skip the overlapping day", "")

    panel.to_csv(CLEAN / "daily_levels.csv")
    daily.to_csv(CLEAN / "daily_changes.csv")
    weekly.to_csv(CLEAN / "weekly_changes.csv")
    monthly.to_csv(CLEAN / "monthly_panel.csv")
    pd.DataFrame(DICT_ROWS, columns=["variable", "description", "unit", "source", "frequency"]) \
        .to_csv(CLEAN / "data_dictionary.csv", index=False)
    pd.DataFrame(QUALITY_LOG).to_csv(CLEAN / "data_quality_log.csv", index=False)

    cov = panel.select_dtypes("number").apply(lambda s: pd.Series(
        {"first": s.first_valid_index(), "last": s.last_valid_index(), "n_obs": s.notna().sum()}))
    cov.T.to_csv(CLEAN / "coverage.csv")
    print(cov.T.to_string())
    print(f"\nrolls detected: " + ", ".join(f"{c}={int(panel[c].sum())}" for c in panel if c.startswith("roll_")))


if __name__ == "__main__":
    main()
